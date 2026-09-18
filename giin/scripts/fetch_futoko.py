#!/usr/bin/env python3
"""不登校に関する公表統計を取り込む。**議員事務所と保護者が「知らない」状態をなくすため。**

文部科学省「児童生徒の問題行動・不登校等生徒指導上の諸課題に関する調査」から取る。
**概要PDFの都道府県別は棒グラフで、数値が図の中に散っている。そこからは読まない。**
同じ調査の統計表が e-Stat に Excel で置いてあるので、そちらを使う。

**表番号は年度でずれる。**（令和5年度の 4-14 が令和6年度は 4-15 になっている）
なので statInfId も表番号も決め打ちせず、**表の名前で検索して新しい年度から取る**。
取ったあとは Excel の見出し（①都道府県別／②指定都市別）で中身を確かめてから使う。
確かめられなければ取り込まない。取れたふりをしない。

    /usr/bin/python3 scripts/fetch_futoko.py          # data/futoko.json と stats.json を更新
    /usr/bin/python3 scripts/fetch_futoko.py --dry    # 取るだけ。保存しない
"""
from __future__ import annotations

import argparse
import datetime
import io
import json
import os
import re
import sys
import urllib.parse
import urllib.request

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "futoko.json")
STATS = os.path.join(ROOT, "data", "stats.json")
UA = {"User-Agent": "Mozilla/5.0 (compatible; xb4g-giin/1.0; +https://xb4g.com/giin/)"}
TOUKEI = "00400304"          # 児童生徒の問題行動・不登校等生徒指導上の諸課題に関する調査
SEARCH = "https://www.e-stat.go.jp/stat-search/files"
DL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={}&fileKind=0"
SURVEY = "文部科学省「児童生徒の問題行動・不登校等生徒指導上の諸課題に関する調査」"
SURVEY_URL = "https://www.e-stat.go.jp/statistics/00400304"
THEME, TRACKER = "kosodate", "futoko"


def get(url: str) -> bytes:
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=180) as r:
        return r.read()


def find_tables(query: str) -> list:
    """表の名前で e-Stat を検索し、(statInfId, 年度) を新しい順に返す。"""
    q = urllib.parse.urlencode({"query": query, "layout": "dataset", "toukei": TOUKEI})
    html = get(f"{SEARCH}?{q}").decode("utf-8", "ignore")
    out, seen = [], set()
    for m in re.finditer(r"file-download\?statInfId=(\d+)", html):
        sid = m.group(1)
        if sid in seen:
            continue
        seen.add(sid)
        ctx = re.sub(r"<[^>]+>", " ", html[max(0, m.start() - 2500): m.end() + 1500])
        y = re.search(r"(20\d\d)年度", ctx)
        out.append((sid, int(y.group(1)) if y else 0))
    out.sort(key=lambda x: -x[1])
    return out


def sheet(sid: str):
    return openpyxl.load_workbook(io.BytesIO(get(DL.format(sid))), data_only=True).active


def cells(ws, r, n=10):
    return [ws.cell(r, c).value for c in range(1, n + 1)]


def head(ws) -> str:
    return " ".join(str(ws.cell(r, 1).value or "") for r in range(1, 4))


def row_by_name(ws, name: str):
    """名前で行を引く。**位置では引かない**（年度で行がずれるため）。"""
    for r in range(1, ws.max_row + 1):
        for c in (1, 2, 3):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() == name:
                return cells(ws, r)
    return None


def num(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return x
    s = re.sub(r"[^\d.\-]", "", str(x))
    if s in ("", "-", "."):
        return None
    return float(s) if "." in s else int(s)


def grab_futoko(rec: dict) -> None:
    """都道府県別・指定都市別 不登校児童生徒数。愛知県と名古屋市だけを持つ。"""
    # **表題には「都道府県別・指定都市別」と両方入る。**（4-15 のような表番号ごとの題）
    # 中身の見分けは、その下にある ①都道府県別（国公私立）／②指定都市別（指定都市立…）。
    # ここを間違えると、県は最新年度・市は前年度、という混ざり方をする。
    def classify(h: str):
        if "②" in h or "指定都市別（指定都市立" in h:
            return "city"
        if "①" in h or "都道府県別（国公私立" in h:
            return "pref"
        return None

    cand = find_tables("都道府県別・指定都市別 不登校児童生徒数")[:8]
    newest = max((y for _, y in cand), default=0)
    got = {}
    for sid, year in cand:
        if year != newest:        # **年度をそろえる。**県と市で年がずれたら比べられない
            continue
        ws = sheet(sid)
        kind = classify(head(ws))
        if kind and kind not in got:
            got[kind] = (ws, year, sid, head(ws))
        if len(got) == 2:
            break

    for kind, want in (("pref", "愛知県"), ("city", "名古屋市")):
        if kind not in got:
            print(f"  不登校({kind}): 表が見つかりませんでした")
            continue
        ws, year, sid, h = got[kind]
        row = row_by_name(ws, want)
        if not row:
            print(f"  不登校({kind}): {want} の行が見つかりませんでした")
            continue
        v = [num(x) for x in row[2:8]]
        if len([x for x in v if x is not None]) < 6:
            print(f"  不登校({kind}): {want} の値がそろいません")
            continue
        d = {"name": want, "year": year,
             "小学校": v[0], "小学校1000人当たり": v[1],
             "中学校": v[2], "中学校1000人当たり": v[3],
             "計": v[4], "計1000人当たり": v[5],
             "source": SURVEY, "source_url": SURVEY_URL, "stat_inf_id": sid}
        # 検算: 小＋中＝計
        if d["小学校"] + d["中学校"] != d["計"]:
            print(f"  不登校({kind}): 小＋中が計と一致しません。取り込みません")
            continue
        rec.setdefault("futoko", {})[kind] = d
        print(f"  不登校({kind}): {want} {year}年度 小{d['小学校']:,}+中{d['中学校']:,}"
              f"={d['計']:,}人（1,000人当たり{d['計1000人当たり']}）")

    # 全国（都道府県表の合計行）
    if "pref" in got:
        ws, year, sid, _ = got["pref"]
        for nm in ("全国", "計", "合計"):
            row = row_by_name(ws, nm)
            if row and num(row[6]):
                rec.setdefault("futoko", {})["all"] = {
                    "name": "全国", "year": year, "計": num(row[6]),
                    "小学校": num(row[2]), "中学校": num(row[4]),
                    "source": SURVEY, "source_url": SURVEY_URL}
                print(f"  不登校(全国): {year}年度 {num(row[6]):,}人")
                break


def grab_center(rec: dict) -> None:
    """教育委員会が設置する「教育支援センター」の状況。愛知県と全国。"""
    for sid, year in find_tables("教育委員会が設置する 教育支援センター の状況")[:5]:
        ws = sheet(sid)
        if "教育支援センター" not in head(ws):
            continue
        a = row_by_name(ws, "愛知県")
        z = row_by_name(ws, "全国")
        if not a:
            continue
        rec["center"] = {
            "year": year, "source": SURVEY, "source_url": SURVEY_URL, "stat_inf_id": sid,
            "愛知県": {"設置数": num(a[2]), "常勤": num(a[3]), "常勤割合": num(a[4]),
                     "非常勤": num(a[5]), "非常勤割合": num(a[6]), "指導員計": num(a[7])},
        }
        if z:
            rec["center"]["全国"] = {"設置数": num(z[2]), "指導員計": num(z[7]),
                                   "非常勤割合": num(z[6])}
        print(f"  教育支援センター: {year}年度 愛知県 {num(a[2])}箇所・"
              f"指導員{num(a[7])}人（非常勤{num(a[6])}%）／全国 "
              f"{num(z[2]) if z else '?'}箇所")
        return
    print("  教育支援センター: 取れませんでした")


# ---- 画面へ ------------------------------------------------------------------
def build_stats(rec: dict) -> list:
    common = {"theme": THEME, "tracker": TRACKER, "auto": "futoko",
              "checked_at": rec.get("checked_at", "")}
    out = []
    f = rec.get("futoko") or {}
    if f.get("pref") or f.get("city"):
        rows, year = [], 0
        for k, lab in (("all", "全国"), ("pref", "愛知県"), ("city", "名古屋市")):
            d = f.get(k)
            if not d:
                continue
            year = max(year, d["year"])
            note = (f"小学校{d['小学校']:,}人／中学校{d['中学校']:,}人"
                    + (f"／1,000人当たり{d['計1000人当たり']}人" if d.get("計1000人当たり") else ""))
            rows.append({"label": lab, "value": f"{d['計']:,}人", "note": note})
        out.append(dict(common, title="不登校の小中学生（都道府県別・指定都市別）",
                        asof=f"令和{year - 2018}年度",
                        rows=rows, source=SURVEY, source_url=SURVEY_URL,
                        caveat="都道府県別は国公私立、指定都市別は指定都市立の小・中学校が対象なので、"
                               "愛知県の数と名古屋市の数はそのまま引き算できません。"
                               "「1,000人当たり」は在籍者数に対する割合なので、規模の違う地域を"
                               "比べるときはこちらを見てください。高校は別の集計です。"))
    c = rec.get("center")
    if c:
        a = c["愛知県"]
        rows = [{"label": "愛知県の設置数", "value": f"{a['設置数']}箇所",
                 "note": "都道府県別には指定都市を含む"},
                {"label": "愛知県の指導員", "value": f"{a['指導員計']}人",
                 "note": f"常勤{a['常勤']}人（{a['常勤割合']}%）／"
                         f"非常勤{a['非常勤']}人（{a['非常勤割合']}%）"}]
        if c.get("全国"):
            z = c["全国"]
            rows.append({"label": "全国の設置数", "value": f"{z['設置数']:,}箇所",
                         "note": f"指導員{z['指導員計']:,}人／非常勤が{z['非常勤割合']}%"})
        out.append(dict(common, title="教育委員会が設置する「教育支援センター」",
                        asof=f"令和{c['year'] - 2018}年度", rows=rows,
                        source=SURVEY, source_url=SURVEY_URL,
                        caveat="市町村や都道府県の教育委員会が設置する、学校の外の通所先です"
                               "（名古屋市では「なごやフレンドリーナウ」）。"
                               "民間のフリースクールはこの数に入りません。"))
    return out


def sync_stats(rec: dict) -> int:
    d = json.load(open(STATS, encoding="utf-8")) if os.path.isfile(STATS) else {"stats": []}
    kept = [s for s in d.get("stats", []) if s.get("auto") != "futoko"]
    made = build_stats(rec)
    d["stats"] = kept + made
    json.dump(d, open(STATS, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return len(made)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    a = ap.parse_args()
    rec = {"checked_at": datetime.datetime.now().strftime("%Y-%m-%d")}
    grab_futoko(rec)
    grab_center(rec)
    if not rec.get("futoko"):
        print("\n！ 何も取れませんでした。保存しません。")
        return 1
    if a.dry:
        print("\n" + json.dumps(rec, ensure_ascii=False, indent=1))
        return 0
    json.dump(rec, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    n = sync_stats(rec)
    print(f"\n{OUT} と data/stats.json（{n}ブロック）を更新しました")
    return 0


if __name__ == "__main__":
    sys.exit(main())
