#!/usr/bin/env python3
"""選挙区ダッシュボードに足す、市区町村ごとの公的な統計を落として data/muni_extra_stats.json にまとめる。

  /usr/bin/python3 scripts/fetch_muni_stats.py          # 落としてあるファイルがあれば使う
  /usr/bin/python3 scripts/fetch_muni_stats.py --fresh  # 落とし直す

**なぜ要るか。** 議員の投稿への返信に「地元の数字カード」を貼るとき、防災と福祉の事業所の数だけでは
話題に合わないことが多かった（2026-10-14 都市農業・子育て・人口減少の投稿）。

入れるもの（どれも国が市区町村ごとに公表している表。数字は書き写すだけで、推計はしない）
  総務省 住民基本台帳に基づく人口（2026年1月1日）  人口・65歳以上・2025年の出生数・2025年の人口の増減
      政令市は区ごとの行がある。コードは6桁（末尾は検査数字）なので頭5桁で引く
  国交省 都市計画現況調査（2024年3月31日）          生産緑地地区の面積（ha）・地区数（都市名で引く）
  こども家庭庁 保育所等関連状況取りまとめ（2025年4月1日）  申込者数・待機児童・いわゆる隠れ待機児童
      隠れ待機児童＝地方単独事業を利用・育児休業中・特定の保育園等のみ希望・求職活動を休止（国の待機児童には数えない人）

元のファイルは outputs/stat_src/ に置く（gitの外）。読むのは表の数字だけ。
"""
from __future__ import annotations

import json
import os
import sys
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "outputs", "stat_src")
OUT = os.path.join(ROOT, "data", "muni_extra_stats.json")

FILES = {
    "nenrei": ("soumu/nenrei_r8.xlsx", "https://www.soumu.go.jp/main_content/000892953.xlsx"),
    "jinko": ("soumu/jinko_r8.xlsx", "https://www.soumu.go.jp/main_content/000892952.xlsx"),
    "seisan": ("mlit/seisan_r6.xls", "https://www.mlit.go.jp/toshi/tosiko/content/001995363.xls"),
    "hoiku": ("cfa/teiin_r7.xlsx", "https://www.cfa.go.jp/assets/contents/node/basic_page/field_ref_resources/"
                                   "b0a8057b-34bf-4c20-84fb-ae592708ca9b/1aa96453/20250828_policies_hoiku_torimatome_r7_04.xlsx"),
}
SOURCES = {
    "jumin": {"name": "総務省 住民基本台帳に基づく人口・人口動態及び世帯数", "as_of": "2026-01-01（人口動態は2025年中）",
              "url": "https://www.soumu.go.jp/main_sosiki/jichi_gyousei/daityo/jinkou_jinkoudoutai-setaisuu.html"},
    "seisan": {"name": "国土交通省 都市計画現況調査（生産緑地地区）", "as_of": "2024-03-31",
               "url": "https://www.mlit.go.jp/toshi/tosiko/toshi_tosiko_tk_000206.html"},
    "hoiku": {"name": "こども家庭庁 保育所等関連状況取りまとめ（申込者の状況）", "as_of": "2025-04-01",
              "url": "https://www.cfa.go.jp/policies/hoiku/torimatome/r7"},
}


def path(k: str) -> str:
    return os.path.join(SRC, FILES[k][0])


def download(fresh: bool) -> None:
    for k, (rel, url) in FILES.items():
        p = os.path.join(SRC, rel)
        if os.path.exists(p) and not fresh:
            continue
        os.makedirs(os.path.dirname(p), exist_ok=True)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (xb4g giin)"})
        with urllib.request.urlopen(req, timeout=120) as r, open(p, "wb") as f:
            f.write(r.read())
        print("落とした", rel)


def num(x) -> int:
    try:
        return int(round(float(x)))
    except (TypeError, ValueError):
        return 0


def jumin() -> dict:
    import openpyxl
    out: dict[str, dict] = {}
    ws = openpyxl.load_workbook(path("nenrei"), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[1]
    total_i = head.index("総数")
    i65 = head.index("65歳～69歳")
    for r in rows:
        code = str(r[0] or "")
        if len(code) != 6 or not code.isdigit() or r[3] != "計" or r[2] in (None, "-"):
            continue
        out.setdefault(code[:5], {})
        out[code[:5]].update({"name": f"{r[1]}{r[2]}", "pop": num(r[total_i]),
                              "pop65": sum(num(v) for v in r[i65:])})
    ws = openpyxl.load_workbook(path("jinko"), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[4]
    ib = head.index("出生者数")
    ic = [i for i, h in enumerate(head) if str(h or "").startswith("増減数")][0]
    for r in rows:
        code = str(r[0] or "")
        if len(code) != 6 or not code.isdigit() or r[2] in (None, "-"):
            continue
        out.setdefault(code[:5], {"name": f"{r[1]}{r[2]}"})
        out[code[:5]].update({"births": num(r[ib]), "change": num(r[ic])})
    return out


def seisan() -> dict:
    import xlrd
    sh = xlrd.open_workbook(path("seisan")).sheets()[0]
    out = {}
    for i in range(sh.nrows):
        pref, city = str(sh.cell_value(i, 0)).strip(), str(sh.cell_value(i, 2)).strip()
        if not pref.endswith(("都", "道", "府", "県")) or not city:
            continue
        ha, n = sh.cell_value(i, 4), sh.cell_value(i, 5)
        if not isinstance(ha, (int, float)):   # 見出しの行（「都道府県」「決定面積」）を飛ばす
            continue
        out[f"{pref}|{city}"] = {"ha10": int(round(float(ha) * 10)), "districts": num(n)}
    return out


def hoiku() -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path("hoiku"), read_only=True, data_only=True)
    ws = [w for w in wb.worksheets if w.title.startswith("申込者")][0]
    rows = list(ws.iter_rows(values_only=True))
    hi = [i for i, r in enumerate(rows) if r and "申込者数" in [str(x or "") for x in r]][0]
    head = [str(x or "") for x in rows[hi]]

    def col(prefix: str) -> int:   # 合計の塊（いちばん左）の列
        return [i for i, h in enumerate(head) if h.startswith(prefix)][0]
    c_app, c_wait = col("申込者数"), col("待機児童")
    c_hidden = [col("地方単独事業"), col("育児休業中"), col("特定の保育園"), col("求職活動を休止")]
    out = {}
    for r in rows[hi + 1:]:
        pref, city = str(r[2] or "").strip(), str(r[3] or "").strip()
        if not pref or not city or not pref.endswith(("都", "道", "府", "県")):
            continue
        out[f"{pref}|{city}"] = {"applicants": num(r[c_app]), "waiting": num(r[c_wait]),
                                 "hidden": sum(num(r[c]) for c in c_hidden)}
    return out


def main() -> int:
    download("--fresh" in sys.argv)
    data = {"sources": SOURCES, "jumin": jumin(), "seisan": seisan(), "hoiku": hoiku()}
    json.dump(data, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    j = data["jumin"]
    print(f"住民基本台帳 {len(j)} 市区町村（人口計 {sum(v.get('pop', 0) for v in j.values()):,}）")
    print(f"生産緑地 {len(data['seisan'])} 都市（{sum(v['ha10'] for v in data['seisan'].values()) / 10:,.1f}ha）")
    h = data["hoiku"]
    print(f"保育 {len(h)} 市区町村（待機児童 {sum(v['waiting'] for v in h.values()):,}人・隠れ {sum(v['hidden'] for v in h.values()):,}人）")
    print("→", OUT, os.path.getsize(OUT) // 1024, "KB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
